"""
Парсер цен: собирает каталог товаров с сайта и сохраняет его в Excel (демо-проект для портфолио).

Сайт books.toscrape.com - учебный магазин, специально созданный для тренировки парсинга.

Что делает:
  1. Проходит по страницам каталога.
  2. С каждой карточки берёт название, цену, рейтинг, наличие и ссылку.
  3. Сохраняет всё в books.xlsx: лист с товарами и лист со сводкой по рейтингам.

Запуск:
  pip install requests beautifulsoup4 openpyxl
  python scraper.py          (первые 5 страниц)
  python scraper.py 50       (все 50 страниц)
"""
import sys
import time
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# ---------- Настройки ----------
START_URL = "https://books.toscrape.com/catalogue/page-1.html"
OUTPUT_FILE = "books.xlsx"
PAUSE = 0.5  # пауза между страницами в секундах, чтобы не нагружать сайт
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}


def parse_page(html: str, page_url: str) -> tuple[list[dict], str | None]:
    """Достаёт товары с одной страницы и ссылку на следующую (или None, если страниц больше нет)."""
    soup = BeautifulSoup(html, "html.parser")
    books = []
    for card in soup.select("article.product_pod"):
        link = card.select_one("h3 a")
        price_text = card.select_one(".price_color").get_text(strip=True)   # например "£51.77"
        rating_word = card.select_one(".star-rating")["class"][1]           # например "Three"
        books.append({
            "title": link["title"],
            "price": float(price_text.replace("£", "").replace("Â", "")),
            "rating": RATINGS.get(rating_word, 0),
            "in_stock": "In stock" in card.select_one(".availability").get_text(),
            "url": urljoin(page_url, link["href"]),
        })
    next_link = soup.select_one("li.next a")
    next_url = urljoin(page_url, next_link["href"]) if next_link else None
    return books, next_url


def scrape(max_pages: int) -> list[dict]:
    """Обходит каталог страница за страницей."""
    all_books, url, page = [], START_URL, 0
    while url and page < max_pages:
        response = requests.get(url, timeout=15)
        response.raise_for_status()          # остановиться с понятной ошибкой, если сайт не ответил
        response.encoding = "utf-8"
        books, url = parse_page(response.text, response.url)
        all_books.extend(books)
        page += 1
        print(f"Страница {page}: найдено {len(books)}, всего {len(all_books)}")
        time.sleep(PAUSE)
    return all_books


def save_excel(books: list[dict], path: str) -> None:
    """Сохраняет товары и сводку в Excel-файл."""
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")

    # Лист 1: все товары
    ws = wb.active
    ws.title = "Books"
    ws.append(["Title", "Price (£)", "Rating", "In stock", "URL"])
    for b in books:
        ws.append([b["title"], b["price"], b["rating"], "Yes" if b["in_stock"] else "No", b["url"]])
    for cell in ws[1]:
        cell.font, cell.fill = header_font, header_fill
    for column, width in zip("ABCDE", (60, 12, 10, 10, 70)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"                   # шапка остаётся на месте при прокрутке
    ws.auto_filter.ref = ws.dimensions       # кнопки фильтра в шапке

    # Лист 2: сводка по рейтингам
    summary = wb.create_sheet("Summary")
    summary.append(["Rating", "Books", "Average price (£)", "Cheapest (£)", "Most expensive (£)"])
    for rating in range(5, 0, -1):
        prices = [b["price"] for b in books if b["rating"] == rating]
        if prices:
            summary.append([rating, len(prices), round(sum(prices) / len(prices), 2), min(prices), max(prices)])
    prices = [b["price"] for b in books]
    summary.append(["All", len(prices), round(sum(prices) / len(prices), 2), min(prices), max(prices)])
    for cell in summary[1]:
        cell.font, cell.fill = header_font, header_fill
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
    for column in "ABCDE":
        summary.column_dimensions[column].width = 22

    wb.save(path)


def main():
    max_pages = int(sys.argv[1]) if len(sys.argv) > 1 else 5
    books = scrape(max_pages)
    if not books:
        raise SystemExit("Ничего не найдено: возможно, изменилась вёрстка сайта.")
    save_excel(books, OUTPUT_FILE)
    print(f"Готово: {len(books)} товаров сохранено в {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
